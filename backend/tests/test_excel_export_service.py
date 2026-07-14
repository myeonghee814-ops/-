import io

from openpyxl import load_workbook

from schemas.analysis import PaperAnalysis
from schemas.comparison import ComparisonResult, ComparisonTableRow
from schemas.search import PaperResult
from services.excel_export_service import build_workbook


def _paper(i: int) -> PaperResult:
    return PaperResult(
        title=f"Paper {i}",
        authors=[f"Author {i}"],
        journal="Journal of Power Sources",
        year=2022,
        citation_count=10 * i,
        doi=f"10.1234/paper.{i}",
        abstract="An abstract about electrolytes. " * 10,
        pdf_url=None,
        published_date="2022-01-01",
        source="semantic_scholar",
    )


def _analysis(i: int) -> PaperAnalysis:
    return PaperAnalysis(
        title=f"Paper {i}",
        authors=[f"Author {i}"],
        journal="Journal of Power Sources",
        electrolyte="1M LiPF6 in EC/DMC",
        salt="LiPF6",
        solvent="EC/DMC",
        additive="FEC",
        cathode="NMC811",
        anode="Graphite",
        separator="Celgard 2325",
        cell_type="Coin cell",
        voltage_window="3.0-4.3 V",
        temperature="25 C",
        formation_protocol="C/20 for 2 cycles",
        cycle_condition="1C/1C, 25 C",
        rate_capability="80% at 5C",
        main_findings=["Improved cycling stability"],
        innovation="Novel additive combination",
        advantages=["Higher capacity retention"],
        limitations=["Limited high-temperature data"],
        future_work=["Test at elevated temperatures"],
    )


def _comparison(paper_count: int) -> ComparisonResult:
    return ComparisonResult(
        paper_count=paper_count,
        common_experimental_conditions=["Most papers use LiPF6-based electrolytes"],
        differences=["Cathode chemistry varies between papers"],
        frequently_used_electrolytes=["1M LiPF6 in EC/DMC"],
        frequently_used_additives=["FEC"],
        most_common_cathode="NMC811",
        most_common_anode="Graphite",
        research_trend="Increasing focus on additive engineering for SEI stability",
        research_gap="Limited high-temperature cycling data",
        potential_future_direction="Systematic study of additive combinations at elevated temperatures",
        comparison_table=[
            ComparisonTableRow(title=f"Paper {i}", electrolyte="1M LiPF6 in EC/DMC", cathode="NMC811")
            for i in range(paper_count)
        ],
    )


def test_build_workbook_creates_four_sheets_in_order() -> None:
    papers = [_paper(i) for i in range(3)]
    analyses = [_analysis(i) for i in range(3)]
    workbook_bytes = build_workbook(papers, analyses, None, "Not enough papers selected.")

    wb = load_workbook(io.BytesIO(workbook_bytes))

    assert wb.sheetnames == ["Summary Table", "Experimental Conditions", "AI Summary", "Comparison"]


def test_summary_sheet_has_expected_header_and_rows() -> None:
    papers = [_paper(i) for i in range(3)]
    analyses = [_analysis(i) for i in range(3)]
    wb = load_workbook(io.BytesIO(build_workbook(papers, analyses, None, "note")))

    ws = wb["Summary Table"]

    assert [cell.value for cell in ws[1]] == [
        "Title",
        "Authors",
        "Journal",
        "Year",
        "Citations",
        "DOI",
        "Published Date",
        "Abstract",
    ]
    assert ws.cell(row=2, column=1).value == "Paper 0"
    assert ws.cell(row=2, column=6).value == "10.1234/paper.0"
    assert ws.max_row == 4  # header + 3 papers


def test_header_row_is_frozen_and_styled() -> None:
    papers = [_paper(0)]
    analyses = [_analysis(0)]
    wb = load_workbook(io.BytesIO(build_workbook(papers, analyses, None, "note")))

    ws = wb["Summary Table"]

    assert ws.freeze_panes == "A2"
    header_cell = ws.cell(row=1, column=1)
    assert header_cell.font.bold is True
    assert header_cell.font.color.rgb == "00FFFFFF"
    assert header_cell.fill.fgColor.rgb == "001E293B"


def test_columns_are_auto_sized() -> None:
    papers = [_paper(0)]
    analyses = [_analysis(0)]
    wb = load_workbook(io.BytesIO(build_workbook(papers, analyses, None, "note")))

    ws = wb["Summary Table"]
    for col_letter in ["A", "B", "C", "D", "E", "F", "G", "H"]:
        assert ws.column_dimensions[col_letter].width and ws.column_dimensions[col_letter].width > 0


def test_experimental_conditions_sheet_handles_missing_analysis() -> None:
    papers = [_paper(0), _paper(1)]
    analyses = [_analysis(0), None]
    wb = load_workbook(io.BytesIO(build_workbook(papers, analyses, None, "note")))

    ws = wb["Experimental Conditions"]

    assert ws.cell(row=2, column=2).value == "1M LiPF6 in EC/DMC"
    assert ws.cell(row=3, column=1).value == "Paper 1"
    assert ws.cell(row=3, column=2).value == "Not analyzed"


def test_ai_summary_sheet_joins_list_fields_as_bullets() -> None:
    papers = [_paper(0)]
    analyses = [_analysis(0)]
    wb = load_workbook(io.BytesIO(build_workbook(papers, analyses, None, "note")))

    ws = wb["AI Summary"]
    headers = [cell.value for cell in ws[1]]

    assert headers == ["Title", "Innovation", "Main Findings", "Advantages", "Limitations", "Future Work"]
    main_findings_col = headers.index("Main Findings") + 1
    assert ws.cell(row=2, column=main_findings_col).value == "• Improved cycling stability"


def test_comparison_sheet_shows_note_when_comparison_unavailable() -> None:
    papers = [_paper(0)]
    analyses = [_analysis(0)]
    wb = load_workbook(io.BytesIO(build_workbook(papers, analyses, None, "Not enough papers selected.")))

    ws = wb["Comparison"]
    values = [cell.value for row in ws.iter_rows() for cell in row if cell.value]

    assert "Not enough papers selected." in values


def test_comparison_sheet_includes_findings_and_table_and_freezes_below_header() -> None:
    papers = [_paper(i) for i in range(10)]
    analyses = [_analysis(i) for i in range(10)]
    comparison = _comparison(10)
    wb = load_workbook(io.BytesIO(build_workbook(papers, analyses, comparison, None)))

    ws = wb["Comparison"]
    values = [cell.value for row in ws.iter_rows() for cell in row if cell.value]

    assert "NMC811" in values
    assert "Common Experimental Conditions" in values
    assert "• Most papers use LiPF6-based electrolytes" in values
    # The per-paper table header should exist and be frozen below (not row 1,
    # since this sheet is a mixed report layout, not a single table).
    assert ws.freeze_panes is not None
    assert ws.freeze_panes != "A2"
