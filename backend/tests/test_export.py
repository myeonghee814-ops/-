from unittest.mock import AsyncMock

import pytest
from httpx import ASGITransport, AsyncClient
from openpyxl import load_workbook
import io

from main import app
from schemas.search import PaperResult


def _sample_paper_payload(i: int) -> dict:
    return PaperResult(
        title=f"Paper {i}",
        authors=[f"Author {i}"],
        journal="Journal of Power Sources",
        year=2022,
        citation_count=i,
        doi=f"10.1234/paper.{i}",
        abstract="An abstract about electrolytes.",
        pdf_url=None,
        published_date="2022-01-01",
        source="semantic_scholar",
    ).model_dump()


@pytest.mark.asyncio
async def test_export_endpoint_returns_xlsx_file(monkeypatch: pytest.MonkeyPatch) -> None:
    fake_workbook_bytes = b"fake-xlsx-bytes"
    # Patched on the route module itself (where the name was imported into),
    # not on services.export_service -- `from x import y` binds a separate
    # reference that patching the origin module wouldn't affect.
    monkeypatch.setattr("api.routes.export.generate_export", AsyncMock(return_value=fake_workbook_bytes))

    transport = ASGITransport(app=app)
    async with AsyncClient(transport=transport, base_url="http://test") as client:
        response = await client.post("/api/v1/export", json={"papers": [_sample_paper_payload(0)]})

    assert response.status_code == 200
    assert response.headers["content-type"] == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    assert "blip-papers-export.xlsx" in response.headers["content-disposition"]
    assert response.content == fake_workbook_bytes


@pytest.mark.asyncio
async def test_export_endpoint_rejects_empty_paper_list() -> None:
    transport = ASGITransport(app=app)
    async with AsyncClient(transport=transport, base_url="http://test") as client:
        response = await client.post("/api/v1/export", json={"papers": []})

    assert response.status_code == 422


@pytest.mark.asyncio
async def test_export_endpoint_returns_a_real_workbook_end_to_end(monkeypatch: pytest.MonkeyPatch) -> None:
    # No mocking of export_service.generate_export here -- only the OpenAI
    # boundary is unavailable (no API key in the test environment), so this
    # exercises the full endpoint -> export_service -> excel_export_service
    # path and confirms a genuinely valid workbook comes back even though AI
    # analysis fails for every paper.
    transport = ASGITransport(app=app)
    async with AsyncClient(transport=transport, base_url="http://test") as client:
        response = await client.post(
            "/api/v1/export", json={"papers": [_sample_paper_payload(0), _sample_paper_payload(1)]}
        )

    assert response.status_code == 200
    wb = load_workbook(io.BytesIO(response.content))
    assert wb.sheetnames == ["Summary Table", "Experimental Conditions", "AI Summary", "Comparison"]
