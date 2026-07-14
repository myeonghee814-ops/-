from types import SimpleNamespace
from unittest.mock import AsyncMock

import pytest
from openpyxl import load_workbook
import io

from schemas.analysis import PaperAnalysis
from schemas.comparison import ComparisonResult
from schemas.search import PaperResult
from services import export_service
from services.external.exceptions import OpenAIAnalysisError


def _paper(i: int, abstract: str | None = "An abstract about electrolytes.") -> PaperResult:
    return PaperResult(
        title=f"Paper {i}",
        authors=[f"Author {i}"],
        journal="Journal of Power Sources",
        year=2022,
        citation_count=i,
        doi=f"10.1234/paper.{i}",
        abstract=abstract,
        pdf_url=None,
        published_date="2022-01-01",
        source="semantic_scholar",
    )


def _analysis(i: int) -> PaperAnalysis:
    return PaperAnalysis(title=f"Paper {i}", cathode="NMC811", anode="Graphite")


def _comparison_result(paper_count: int) -> ComparisonResult:
    return ComparisonResult(paper_count=paper_count, most_common_cathode="NMC811")


@pytest.mark.asyncio
async def test_analyze_all_skips_papers_without_abstract(monkeypatch: pytest.MonkeyPatch) -> None:
    mock_analyze = AsyncMock(return_value=_analysis(0))
    monkeypatch.setattr(export_service, "analyze_paper", mock_analyze)

    papers = [_paper(0, abstract=None), _paper(1)]
    results = await export_service._analyze_all(papers)

    assert results[0] is None
    assert results[1] is not None
    mock_analyze.assert_awaited_once()


@pytest.mark.asyncio
async def test_analyze_all_tolerates_individual_failures(monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr(export_service, "analyze_paper", AsyncMock(side_effect=OpenAIAnalysisError("boom")))

    results = await export_service._analyze_all([_paper(0), _paper(1)])

    assert results == [None, None]


@pytest.mark.asyncio
async def test_compare_if_possible_returns_note_below_minimum(monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr(export_service, "get_settings", lambda: SimpleNamespace(COMPARISON_MIN_PAPERS=10))

    comparison, note = await export_service._compare_if_possible([_analysis(0)], total_papers=1)

    assert comparison is None
    assert "at least 10" in note
    assert "1 of 1" in note


@pytest.mark.asyncio
async def test_compare_if_possible_runs_comparison_at_minimum(monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr(export_service, "get_settings", lambda: SimpleNamespace(COMPARISON_MIN_PAPERS=2))
    monkeypatch.setattr(export_service, "compare_papers", AsyncMock(return_value=_comparison_result(2)))

    analyses = [_analysis(0), _analysis(1)]
    comparison, note = await export_service._compare_if_possible(analyses, total_papers=2)

    assert note is None
    assert comparison.most_common_cathode == "NMC811"


@pytest.mark.asyncio
async def test_compare_if_possible_degrades_gracefully_on_provider_failure(monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr(export_service, "get_settings", lambda: SimpleNamespace(COMPARISON_MIN_PAPERS=2))
    monkeypatch.setattr(export_service, "compare_papers", AsyncMock(side_effect=OpenAIAnalysisError("boom")))

    comparison, note = await export_service._compare_if_possible([_analysis(0), _analysis(1)], total_papers=2)

    assert comparison is None
    assert note == "AI comparison is currently unavailable."


@pytest.mark.asyncio
async def test_generate_export_produces_a_valid_workbook_even_without_ai(monkeypatch: pytest.MonkeyPatch) -> None:
    # No AI available at all (e.g. OPENAI_API_KEY unset) -- export should still succeed.
    monkeypatch.setattr(export_service, "analyze_paper", AsyncMock(side_effect=OpenAIAnalysisError("no key")))
    monkeypatch.setattr(export_service, "get_settings", lambda: SimpleNamespace(COMPARISON_MIN_PAPERS=10))

    papers = [_paper(i) for i in range(3)]
    workbook_bytes = await export_service.generate_export(papers)

    wb = load_workbook(io.BytesIO(workbook_bytes))
    assert wb.sheetnames == ["Summary Table", "Experimental Conditions", "AI Summary", "Comparison"]
    assert wb["Summary Table"].cell(row=2, column=1).value == "Paper 0"
    assert wb["Experimental Conditions"].cell(row=2, column=2).value == "Not analyzed"
