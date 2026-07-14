"""Builds the exported .xlsx workbook from already-fetched data.

Pure and synchronous by design: no network calls, no knowledge of AI
providers. services/export_service.py gathers papers/analyses/comparison
data (which does require network calls) and hands them to
`build_workbook` here, which only ever touches openpyxl.
"""

import io

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.worksheet import Worksheet

from schemas.analysis import PaperAnalysis
from schemas.comparison import ComparisonResult
from schemas.search import PaperResult

_HEADER_FILL = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
_HEADER_FONT = Font(color="FFFFFF", bold=True, size=11)
_HEADER_ALIGNMENT = Alignment(horizontal="center", vertical="center", wrap_text=True)
_BAND_FILL = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")
_BORDER_SIDE = Side(style="thin", color="CBD5E1")
_CELL_BORDER = Border(left=_BORDER_SIDE, right=_BORDER_SIDE, top=_BORDER_SIDE, bottom=_BORDER_SIDE)
_WRAP_TOP = Alignment(wrap_text=True, vertical="top")
_TITLE_FONT = Font(bold=True, size=14, color="1E293B")
_SECTION_FONT = Font(bold=True, size=12, color="1E293B")
_NOTE_FONT = Font(italic=True, color="64748B")

_EXPERIMENTAL_FIELDS = [
    ("Electrolyte", "electrolyte"),
    ("Salt", "salt"),
    ("Solvent", "solvent"),
    ("Additive", "additive"),
    ("Cathode", "cathode"),
    ("Anode", "anode"),
    ("Separator", "separator"),
    ("Cell Type", "cell_type"),
    ("Voltage Window", "voltage_window"),
    ("Temperature", "temperature"),
    ("Formation Protocol", "formation_protocol"),
    ("Cycle Condition", "cycle_condition"),
    ("Rate Capability", "rate_capability"),
]

_LIST_FIELDS = [
    ("Main Findings", "main_findings"),
    ("Advantages", "advantages"),
    ("Limitations", "limitations"),
    ("Future Work", "future_work"),
]

_COMPARISON_TABLE_HEADERS = [
    "Title",
    "Electrolyte",
    "Salt",
    "Additive",
    "Cathode",
    "Anode",
    "Separator",
    "Cell Type",
    "Voltage Window",
    "Temperature",
    "Cycle Condition",
    "Main Finding",
    "Advantages",
    "Limitations",
]


def build_workbook(
    papers: list[PaperResult],
    analyses: list[PaperAnalysis | None],
    comparison: ComparisonResult | None,
    comparison_note: str | None,
) -> bytes:
    """Build the 4-sheet export workbook and return its raw .xlsx bytes."""
    wb = Workbook()
    wb.remove(wb.active)

    _build_summary_sheet(wb, papers, analyses)
    _build_experimental_conditions_sheet(wb, papers, analyses)
    _build_ai_summary_sheet(wb, papers, analyses)
    _build_comparison_sheet(wb, comparison, comparison_note)

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def _build_summary_sheet(wb: Workbook, papers: list[PaperResult], analyses: list[PaperAnalysis | None]) -> None:
    ws = wb.create_sheet("Paper Summary")
    headers = [
        "Title",
        "Authors",
        "Journal",
        "Year",
        "Citations",
        "Battery System",
        "Electrolyte",
        "Main Contribution",
        "DOI",
        "Published Date",
        "Abstract",
    ]
    _write_header_row(ws, headers)

    for row_idx, (paper, analysis) in enumerate(zip(papers, analyses), start=2):
        values = [
            paper.title,
            ", ".join(paper.authors) or "—",
            paper.journal or "—",
            paper.year or "—",
            paper.citation_count if paper.citation_count is not None else 0,
            (analysis.battery_system if analysis else None) or "—",
            (analysis.electrolyte if analysis else None) or "—",
            (analysis.innovation if analysis else None) or "—",
            paper.doi or "—",
            paper.published_date or "—",
            paper.abstract or "—",
        ]
        _write_data_row(ws, row_idx, values)

    _apply_zebra_stripes(ws, len(papers))
    _auto_size_columns(ws)


def _build_experimental_conditions_sheet(
    wb: Workbook, papers: list[PaperResult], analyses: list[PaperAnalysis | None]
) -> None:
    ws = wb.create_sheet("Experimental Conditions")
    headers = ["Title"] + [label for label, _ in _EXPERIMENTAL_FIELDS]
    _write_header_row(ws, headers)

    for row_idx, (paper, analysis) in enumerate(zip(papers, analyses), start=2):
        if analysis is None:
            values = [paper.title] + ["Not analyzed"] * len(_EXPERIMENTAL_FIELDS)
        else:
            values = [paper.title] + [getattr(analysis, field) or "—" for _, field in _EXPERIMENTAL_FIELDS]
        _write_data_row(ws, row_idx, values)

    _apply_zebra_stripes(ws, len(papers))
    _auto_size_columns(ws)


def _build_ai_summary_sheet(wb: Workbook, papers: list[PaperResult], analyses: list[PaperAnalysis | None]) -> None:
    ws = wb.create_sheet("AI Summary")
    headers = ["Title", "Innovation"] + [label for label, _ in _LIST_FIELDS]
    _write_header_row(ws, headers)

    for row_idx, (paper, analysis) in enumerate(zip(papers, analyses), start=2):
        if analysis is None:
            values = [paper.title, "Not analyzed yet"] + ["—"] * len(_LIST_FIELDS)
        else:
            values = [paper.title, analysis.innovation or "—"] + [
                _join_bullets(getattr(analysis, field)) for _, field in _LIST_FIELDS
            ]
        _write_data_row(ws, row_idx, values)
        ws.row_dimensions[row_idx].height = 90

    _apply_zebra_stripes(ws, len(papers))
    _auto_size_columns(ws, max_width=50)


def _build_comparison_sheet(wb: Workbook, comparison: ComparisonResult | None, comparison_note: str | None) -> None:
    ws = wb.create_sheet("Comparison")

    row = 1
    ws.cell(row=row, column=1, value="AI Comparison").font = _TITLE_FONT
    row += 2

    if comparison is None:
        note_cell = ws.cell(row=row, column=1, value=comparison_note or "Comparison unavailable.")
        note_cell.font = _NOTE_FONT
        note_cell.alignment = _WRAP_TOP
        _auto_size_columns(ws, max_width=80, min_width=14)
        return

    row = _write_label_value(ws, row, "Papers Compared", str(comparison.paper_count))
    row = _write_label_value(ws, row, "Most Common Cathode", comparison.most_common_cathode or "—")
    row = _write_label_value(ws, row, "Most Common Anode", comparison.most_common_anode or "—")
    row = _write_label_value(ws, row, "Research Trend", comparison.research_trend or "—")
    row = _write_label_value(ws, row, "Research Gap", comparison.research_gap or "—")
    row = _write_label_value(ws, row, "Potential Future Direction", comparison.potential_future_direction or "—")
    row += 1

    row = _write_bullet_section(ws, row, "Common Experimental Conditions", comparison.common_experimental_conditions)
    row = _write_bullet_section(ws, row, "Differences", comparison.differences)
    row = _write_bullet_section(ws, row, "Frequently Used Electrolytes", comparison.frequently_used_electrolytes)
    row = _write_bullet_section(ws, row, "Frequently Used Additives", comparison.frequently_used_additives)
    row += 1

    table_header_row = row
    _write_header_row(ws, _COMPARISON_TABLE_HEADERS, header_row=table_header_row)

    for offset, table_row in enumerate(comparison.comparison_table):
        data_row = table_header_row + 1 + offset
        values = [
            table_row.title,
            table_row.electrolyte,
            table_row.salt,
            table_row.additive,
            table_row.cathode,
            table_row.anode,
            table_row.separator,
            table_row.cell_type,
            table_row.voltage_window,
            table_row.temperature,
            table_row.cycle_condition,
            table_row.main_finding,
            table_row.advantages,
            table_row.limitations,
        ]
        _write_data_row(ws, data_row, [v or "—" for v in values])
        if offset % 2 == 1:
            for col_idx in range(1, len(_COMPARISON_TABLE_HEADERS) + 1):
                ws.cell(row=data_row, column=col_idx).fill = _BAND_FILL

    # Freeze below the table header rather than row 1, since that's the
    # part of this "report" sheet with enough rows to need scrolling.
    ws.freeze_panes = f"A{table_header_row + 1}"
    _auto_size_columns(ws, max_width=60, min_width=14)


def _write_header_row(ws: Worksheet, headers: list[str], header_row: int = 1) -> None:
    for col_idx, header in enumerate(headers, start=1):
        cell = ws.cell(row=header_row, column=col_idx, value=header)
        cell.font = _HEADER_FONT
        cell.fill = _HEADER_FILL
        cell.alignment = _HEADER_ALIGNMENT
        cell.border = _CELL_BORDER
    ws.row_dimensions[header_row].height = 26
    if header_row == 1:
        ws.freeze_panes = "A2"


def _write_data_row(ws: Worksheet, row_idx: int, values: list) -> None:
    for col_idx, value in enumerate(values, start=1):
        cell = ws.cell(row=row_idx, column=col_idx, value=value)
        cell.alignment = _WRAP_TOP
        cell.border = _CELL_BORDER


def _write_label_value(ws: Worksheet, row: int, label: str, value: str) -> int:
    label_cell = ws.cell(row=row, column=1, value=label)
    label_cell.font = _SECTION_FONT
    label_cell.alignment = Alignment(vertical="top")
    value_cell = ws.cell(row=row, column=2, value=value)
    value_cell.alignment = _WRAP_TOP
    return row + 1


def _write_bullet_section(ws: Worksheet, row: int, label: str, items: list[str]) -> int:
    ws.cell(row=row, column=1, value=label).font = _SECTION_FONT
    row += 1
    if not items:
        ws.cell(row=row, column=1, value="—").alignment = _WRAP_TOP
        return row + 1
    for item in items:
        cell = ws.cell(row=row, column=1, value=f"• {item}")
        cell.alignment = _WRAP_TOP
        row += 1
    return row


def _join_bullets(items: list[str]) -> str:
    return "\n".join(f"• {item}" for item in items) if items else "—"


def _apply_zebra_stripes(ws: Worksheet, num_data_rows: int, start_row: int = 2) -> None:
    if num_data_rows <= 0:
        return
    max_col = ws.max_column
    for offset in range(num_data_rows):
        if offset % 2 == 1:
            row = start_row + offset
            for col_idx in range(1, max_col + 1):
                ws.cell(row=row, column=col_idx).fill = _BAND_FILL


def _auto_size_columns(ws: Worksheet, max_width: int = 60, min_width: int = 10) -> None:
    for column_cells in ws.columns:
        # Multi-line cells (bullet lists) shouldn't force the column as wide
        # as their longest line; base width on the longest single line instead.
        lines = [line for cell in column_cells if cell.value is not None for line in str(cell.value).split("\n")]
        if not lines:
            continue
        longest_line = max(len(line) for line in lines)
        width = max(min_width, min(longest_line + 2, max_width))
        ws.column_dimensions[column_cells[0].column_letter].width = width
